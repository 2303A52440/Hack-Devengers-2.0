"""
Script to generate OmniFlow AI Pitch Deck PowerPoint (.pptx)
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    # 16:9 Widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    DARK_BG = RGBColor(15, 23, 42)      # slate-900
    CARD_BG = RGBColor(30, 41, 59)      # slate-800
    TEXT_WHITE = RGBColor(255, 255, 255)
    TEXT_MUTED = RGBColor(148, 163, 184) # slate-400
    EMERALD = RGBColor(16, 185, 129)    # emerald-500
    SKY_BLUE = RGBColor(56, 189, 248)   # sky-400

    blank_layout = prs.slide_layouts[6]

    def add_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = DARK_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="HACK DEVENGERS 2.0 PROJECT PITCH"):
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.2))
        tf = tx_box.text_frame
        tf.word_wrap = True
        
        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = EMERALD
        p0.space_after = Pt(4)

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(28)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_WHITE

    # ----------------------------------------------------
    # SLIDE 1: TITLE SLIDE
    # ----------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    add_slide_background(slide1)

    tbox = slide1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(3.5))
    tf1 = tbox.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "OmniFlow AI"
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    p2 = tf1.add_paragraph()
    p2.text = "Autonomous Multimodal WhatsApp CRM & Real-Time UPI Billing Engine"
    p2.font.size = Pt(24)
    p2.font.color.rgb = TEXT_WHITE
    p2.space_before = Pt(10)
    p2.space_after = Pt(30)

    p3 = tf1.add_paragraph()
    p3.text = "Submitted by: Seetharam Ravula | Hack Devengers 2.0"
    p3.font.size = Pt(14)
    p3.font.color.rgb = SKY_BLUE

    # ----------------------------------------------------
    # SLIDE 2: THE PROBLEM
    # ----------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_slide_background(slide2)
    add_header(slide2, "The Real-World Problem Facing Indian MSMEs")

    problems = [
        ("🔴 High Lead Drop-off Rates", "Inquiries on WhatsApp go unaddressed for hours because manual customer support reps cannot respond fast enough, causing potential buyers to switch to competitors."),
        ("🔴 Payment Friction & Abandonment", "Sharing bank details or generic payment links manually creates trust issues and delays, leading customers to abandon purchase decisions before checkout."),
        ("🔴 Fragmented Lead Tracking", "Customer chat histories, unpaid Khatas, and transaction receipts stay trapped inside personal phone screens without a unified CRM or payment reconciliation tool.")
    ]

    for i, (p_title, p_desc) in enumerate(problems):
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i*4.0), Inches(2.0), Inches(3.7), Inches(4.5))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = RGBColor(51, 65, 85)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = p_title
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(14)

        p_body = tf.add_paragraph()
        p_body.text = p_desc
        p_body.font.size = Pt(13)
        p_body.font.color.rgb = TEXT_MUTED

    # ----------------------------------------------------
    # SLIDE 3: THE SOLUTION
    # ----------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_slide_background(slide3)
    add_header(slide3, "OmniFlow AI: Autonomous Sales & Instant Billing Engine")

    solutions = [
        ("🎙️ Multimodal & Multilingual AI", "Processes both text messages and voice notes in English, Hindi, Hinglish, and regional code-switching languages with instant response generation."),
        ("⚡ NPCI-Compliant UPI QR Billing", "Automatically generates dynamic GST invoices with scannable UPI QR images and instant deep links (GPay, PhonePe, Paytm, BHIM)."),
        ("📊 Real-Time Web CRM Dashboard", "Sleek dark-mode dashboard displaying lead statuses (NEW, INVOICED, CONVERTED), sentiment analytics, and one-click payment reminders.")
    ]

    for i, (s_title, s_desc) in enumerate(solutions):
        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i*4.0), Inches(2.0), Inches(3.7), Inches(4.5))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = EMERALD

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = s_title
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = EMERALD
        p.space_after = Pt(14)

        p_body = tf.add_paragraph()
        p_body.text = s_desc
        p_body.font.size = Pt(13)
        p_body.font.color.rgb = TEXT_WHITE

    # ----------------------------------------------------
    # SLIDE 4: SYSTEM ARCHITECTURE & TECH STACK
    # ----------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_slide_background(slide4)
    add_header(slide4, "Technical Architecture & Component Breakdown")

    stack_items = [
        ("Backend Server", "Python 3.10, FastAPI, Uvicorn (High performance async REST API)"),
        ("Database Layer", "SQLite, SQLAlchemy ORM (Lightweight lead & invoice tracking)"),
        ("AI & Voice Engine", "VoiceAudioProcessor, Keyword Intent Router, Sentiment Scoring"),
        ("Payment Generator", "NPCI UPI URI Spec, qrcode matrix generator, Pillow (PIL)"),
        ("Admin Dashboard", "HTML5, Tailwind CSS, Alpine.js, Live Web Sandbox Simulator")
    ]

    for i, (tech_name, tech_detail) in enumerate(stack_items):
        y_pos = Inches(2.0 + i*1.0)
        box = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), y_pos, Inches(11.7), Inches(0.8))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = SKY_BLUE

        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.15)

        p = tf.paragraphs[0]
        p.text = f"{tech_name}: "
        p.font.bold = True
        p.font.size = Pt(15)
        p.font.color.rgb = SKY_BLUE

        run = p.add_run()
        run.text = tech_detail
        run.font.bold = False
        run.font.color.rgb = TEXT_WHITE

    # ----------------------------------------------------
    # SLIDE 5: UNIQUENESS & INNOVATION
    # ----------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    add_slide_background(slide5)
    add_header(slide5, "Why OmniFlow AI Stands Out (Innovation Factor)")

    innovations = [
        "✨ Multilingual Voice Note Speech Parser: Solves the literacy and speed barrier for Indian buyers by parsing voice recordings into structured queries.",
        "✨ Zero-Friction UPI Checkout: Eliminates typing bank account numbers; buyers simply scan the auto-generated QR code or tap the UPI payment link.",
        "✨ Automated Status Sync & Reminders: Automatically converts lead status upon payment detection and dispatches single-click payment reminders.",
        "✨ Modular India-First Design: Tailored for Indian coaching institutes, retail, and local service vendors."
    ]

    for i, inn_text in enumerate(innovations):
        y_pos = Inches(2.0 + i*1.2)
        box = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y_pos, Inches(11.7), Inches(1.0))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.fill.background()

        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.2)

        p = tf.paragraphs[0]
        p.text = inn_text
        p.font.size = Pt(15)
        p.font.color.rgb = TEXT_WHITE

    # ----------------------------------------------------
    # SLIDE 6: LINKS & DEMO
    # ----------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    add_slide_background(slide6)
    add_header(slide6, "Project Resources & Links")

    tbox = slide6.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(4.5))
    tf6 = tbox.text_frame
    tf6.word_wrap = True

    p = tf6.paragraphs[0]
    p.text = "🐙 GitHub Source Code Repository:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    p_sub1 = tf6.add_paragraph()
    p_sub1.text = "https://github.com/2303A52440/Hack-Devengers-2.0.git"
    p_sub1.font.size = Pt(16)
    p_sub1.font.color.rgb = SKY_BLUE
    p_sub1.space_after = Pt(20)

    p2 = tf6.add_paragraph()
    p2.text = "🌐 Live Public Web Deployment:"
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = EMERALD

    p_sub2 = tf6.add_paragraph()
    p_sub2.text = "https://shaky-clubs-pump.loca.lt/dashboard"
    p_sub2.font.size = Pt(16)
    p_sub2.font.color.rgb = SKY_BLUE
    p_sub2.space_after = Pt(30)

    p3 = tf6.add_paragraph()
    p3.text = "Thank You! Hack Devengers 2.0 | Build • Learn • Lead • Impact"
    p3.font.size = Pt(20)
    p3.font.bold = True
    p3.font.color.rgb = TEXT_WHITE

    output_path = os.path.join(os.path.dirname(__file__), "OmniFlow_AI_Pitch_Deck.pptx")
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    build_presentation()
